from fastapi import APIRouter, UploadFile, File, Form, HTTPException, Depends
from sqlalchemy.orm import Session
from typing import List
from database import get_db
from models import Resume, User
from routers.auth import get_current_user
import fitz  # PyMuPDF
import spacy
import re
from io import BytesIO
from docx import Document
from fpdf import FPDF
from fastapi.responses import StreamingResponse, JSONResponse
import tempfile
import os

router = APIRouter(prefix="/analytics", tags=["AI Matching"])
nlp = spacy.load("en_core_web_sm")

# ---------------------------
# Helper Functions
# ---------------------------
def extract_text_from_pdf(file: UploadFile) -> str:
    with fitz.open(stream=file.file.read(), filetype="pdf") as doc:
        return "\n".join([page.get_text() for page in doc])

def extract_text_from_docx(file: UploadFile) -> str:
    doc = Document(BytesIO(file.file.read()))
    return "\n".join([p.text for p in doc.paragraphs])

def extract_top_skills(text: str) -> List[str]:
    skills_block = ""
    for match in re.finditer(r"(?i)(skills?)[\s:\n]+(.+?)(?=\n\S|$)", text, re.DOTALL):
        skills_block += match.group(2) + "\n"
    words = set()
    for line in skills_block.splitlines():
        clean_line = re.sub(r"^[•\-–\*\s]+", "", line)
        words.update([w.strip() for w in re.split(r",|\||\s{2,}", clean_line) if w.strip()])
    return list(words)

def match_skills(resume_skills: List[str], required_skills: List[str]) -> int:
    matched = [skill for skill in resume_skills if skill.lower() in [s.lower() for s in required_skills]]
    return int((len(matched) / len(required_skills)) * 100) if required_skills else 0

def match_education(resume_text: str, preferred_education: str) -> int:
    return int(preferred_education.lower() in resume_text.lower()) * 100 if preferred_education else 0

def match_experience(resume_text: str, required_years: int) -> int:
    matches = re.findall(r"(\d+)\+?\s*(?:years?|yrs?)", resume_text.lower())
    alt_match = re.findall(r"over\s+(\d+)\s*(?:years?|yrs?)", resume_text.lower())
    years = max([int(m) for m in matches + alt_match], default=0)
    return 100 if years >= required_years else int((years / required_years) * 100) if required_years else 0

# ---------------------------
# PDF Generation
# ---------------------------
def generate_pdf_report(results: List[dict]) -> bytes:
    pdf = FPDF()
    pdf.add_page()

    font_path = os.path.join("fonts", "DejaVuSans.ttf")
    if not os.path.exists(font_path):
        raise HTTPException(status_code=500, detail="Font file not found at 'fonts/DejaVuSans.ttf'")

    pdf.add_font("DejaVu", "", font_path, uni=True)
    pdf.set_font("DejaVu", size=14)
    pdf.cell(200, 10, "Resume Match Report", ln=True, align="C")

    pdf.set_font("DejaVu", size=12)
    pdf.ln(10)
    for item in results:
        pdf.multi_cell(0, 10, f"Resume: {item['filename']}")
        pdf.multi_cell(0, 10, f"Match Score: {item['score']}%")
        pdf.multi_cell(0, 10, f"Top Skills: {', '.join(item['top_skills']) if item['top_skills'] else 'None detected'}")
        pdf.ln(3)

    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp_file:
        pdf.output(tmp_file.name)
        tmp_file.seek(0)
        return tmp_file.read()

# ---------------------------
# PDF-Only Endpoint
# ---------------------------
@router.post("/match")
async def match_resumes_to_pdf(
    skills: str = Form(...),
    experience: str = Form(...),
    education: str = Form(...),
    resumes: List[UploadFile] = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    results = await compute_resume_match(skills, experience, education, resumes)
    pdf_bytes = generate_pdf_report(results)
    return StreamingResponse(BytesIO(pdf_bytes), media_type="application/pdf", headers={
        "Content-Disposition": "attachment; filename=resume_match_report.pdf"
    })

# ---------------------------
# JSON Results Endpoint for Chart
# ---------------------------
@router.post("/match-full")
async def match_resumes_json(
    skills: str = Form(...),
    experience: str = Form(...),
    education: str = Form(...),
    resumes: List[UploadFile] = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    results = await compute_resume_match(skills, experience, education, resumes)
    return JSONResponse(content={"matches": results})

# ---------------------------
# Reusable Match Logic
# ---------------------------
async def compute_resume_match(skills, experience, education, resumes):
    results = []
    skill_list = [s.strip() for s in skills.split(",") if s.strip()]
    required_years = int(experience) if experience.isdigit() else 0

    for resume_file in resumes:
        file_ext = resume_file.filename.lower()
        resume_file.file.seek(0)
        try:
            if file_ext.endswith(".pdf"):
                resume_text = extract_text_from_pdf(resume_file)
            elif file_ext.endswith(".docx"):
                resume_text = extract_text_from_docx(resume_file)
            else:
                raise HTTPException(status_code=400, detail=f"Unsupported file type: {resume_file.filename}")

            top_skills = extract_top_skills(resume_text)
            score_skills = match_skills(top_skills, skill_list)
            score_edu = match_education(resume_text, education)
            score_exp = match_experience(resume_text, required_years)
            final_score = int((score_skills + score_edu + score_exp) / 3)

            results.append({
                "filename": resume_file.filename,
                "score": final_score,
                "top_skills": top_skills
            })

        except Exception as e:
            raise HTTPException(status_code=500, detail=f"Error processing {resume_file.filename}: {str(e)}")

    results.sort(key=lambda x: x["score"], reverse=True)
    return results
